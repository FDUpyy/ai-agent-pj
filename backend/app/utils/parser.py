import logging
import os
import tempfile
import traceback
from typing import List, Optional

import fitz
import openpyxl
from docx import Document as DocxDocument
from langchain_core.documents import Document
from pptx import Presentation

# Configure logging
logging.basicConfig(level=logging.ERROR)
logger = logging.getLogger(__name__)


class UniversalParser:
    def __init__(self) -> None:
        self.whisper_model = None
        self.dashscope_client = None
        self.dashscope_api_key = os.getenv("DASHSCOPE_API_KEY", "")

    def parse_file(self, file_path: str) -> List[Document]:
        return self.parse(file_path)

    def parse(self, file_path: str) -> List[Document]:
        try:
            ext = os.path.splitext(file_path)[1].lower()
            if ext == ".pdf":
                return self._parse_pdf(file_path)
            if ext in {".ppt", ".pptx"}:
                return self._parse_ppt(file_path)
            if ext in {".xlsx", ".xls"}:
                return self._parse_excel(file_path)
            if ext in {".docx"}:
                return self._parse_docx(file_path)
            if ext in {".mp3", ".wav", ".m4a", ".flac"}:
                return self._parse_audio(file_path)
            if ext in {".mp4", ".mov", ".avi", ".mkv"}:
                return self._parse_video(file_path)
            if ext in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
                return self._parse_image(file_path)
            if ext == ".txt":
                return self._parse_txt(file_path)
            
            source = os.path.abspath(file_path)
            message = "暂不支持此格式，请转换为 .docx 或 .pdf 上传"
            return [
                Document(
                    page_content=message,
                    metadata={"source": source, "unsupported": True},
                )
            ]
        except Exception as e:
            logger.error(f"Global parsing error for {file_path}: {str(e)}")
            logger.error(traceback.format_exc())
            source = os.path.abspath(file_path)
            return [
                Document(
                    page_content=f"[文件解析失败: {str(e)}]",
                    metadata={"source": source, "error": True},
                )
            ]

    def _parse_pdf(self, file_path: str) -> List[Document]:
        try:
            source = os.path.abspath(file_path)
            docs: List[Document] = []
            pdf = fitz.open(file_path)
            for index, page in enumerate(pdf):
                text = page.get_text("text") or ""
                text = text.strip()
                if not text:
                    continue
                docs.append(
                    Document(
                        page_content=text,
                        metadata={"source": source, "page": index + 1},
                    )
                )
            pdf.close()
            return docs
        except Exception as e:
            logger.error(f"PDF parsing error: {str(e)}")
            return [
                Document(
                    page_content=f"[PDF解析失败: {str(e)}]",
                    metadata={"source": os.path.abspath(file_path), "error": True},
                )
            ]

    def _parse_ppt(self, file_path: str) -> List[Document]:
        try:
            source = os.path.abspath(file_path)
            docs: List[Document] = []
            presentation = Presentation(file_path)
            for index, slide in enumerate(presentation.slides):
                parts = []
                for shape in slide.shapes:
                    if hasattr(shape, "text"):
                        value = shape.text or ""
                        if value.strip():
                            parts.append(value)
                text = "\n".join(parts).strip()
                if not text:
                    continue
                docs.append(
                    Document(
                        page_content=text,
                        metadata={"source": source, "page": index + 1},
                    )
                )
            return docs
        except Exception as e:
            logger.error(f"PPT parsing error: {str(e)}")
            return [
                Document(
                    page_content=f"[PPT解析失败: {str(e)}]",
                    metadata={"source": os.path.abspath(file_path), "error": True},
                )
            ]

    def _parse_excel(self, file_path: str) -> List[Document]:
        try:
            source = os.path.abspath(file_path)
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            lines: List[str] = []
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells = [str(cell) for cell in row if cell is not None]
                    if cells:
                        lines.append(" ".join(cells))
            text = "\n".join(lines).strip()
            if not text:
                return []
            return [
                Document(
                    page_content=text,
                    metadata={"source": source},
                )
            ]
        except Exception as e:
            logger.error(f"Excel parsing error: {str(e)}")
            return [
                Document(
                    page_content=f"[Excel解析失败: {str(e)}]",
                    metadata={"source": os.path.abspath(file_path), "error": True},
                )
            ]

    def _parse_docx(self, file_path: str) -> List[Document]:
        try:
            doc = DocxDocument(file_path)
            text = "\n".join(p.text or "" for p in doc.paragraphs)
            text = text.strip()
            if not text:
                return []
            source = os.path.abspath(file_path)
            return [
                Document(
                    page_content=text,
                    metadata={"source": source},
                )
            ]
        except Exception as e:
            logger.error(f"Docx parsing error: {str(e)}")
            return [
                Document(
                    page_content=f"[Word解析失败: {str(e)}]",
                    metadata={"source": os.path.abspath(file_path), "error": True},
                )
            ]

    def _ensure_whisper_model(self) -> None:
        if self.whisper_model is not None:
            return
        import whisper

        print("Loading Whisper base model for audio transcription, this may take a while...")
        self.whisper_model = whisper.load_model("base")

    def _parse_audio(self, file_path: str) -> List[Document]:
        try:
            self._ensure_whisper_model()
            print(f"Transcribing audio file with Whisper: {file_path}")
            result = self.whisper_model.transcribe(file_path)
            source = os.path.abspath(file_path)
            docs: List[Document] = []
            segments = result.get("segments") or []
            for segment in segments:
                text = (segment.get("text") or "").strip()
                if not text:
                    continue
                start = segment.get("start")
                end = segment.get("end")
                docs.append(
                    Document(
                        page_content=text,
                        metadata={"source": source, "start": start, "end": end},
                    )
                )
            if docs:
                return docs
            text = (result.get("text") or "").strip()
            if not text:
                return []
            return [
                Document(
                    page_content=text,
                    metadata={"source": source},
                )
            ]
        except Exception as e:
            logger.error(f"Audio parsing error: {str(e)}")
            return [
                Document(
                    page_content=f"[音频解析失败: {str(e)}]",
                    metadata={"source": os.path.abspath(file_path), "error": True},
                )
            ]

    def _parse_video(self, file_path: str) -> List[Document]:
        try:
            from moviepy.editor import VideoFileClip
            
            clip = VideoFileClip(file_path)
            temp_file: Optional[str] = None
            try:
                with tempfile.NamedTemporaryFile(
                    suffix=".mp3", delete=False
                ) as temp:
                    temp_file = temp.name
                clip.audio.write_audiofile(temp_file, verbose=False, logger=None)
                docs = self._parse_audio(temp_file)
            finally:
                clip.close()
                if temp_file and os.path.exists(temp_file):
                    os.remove(temp_file)
            return docs
        except ImportError:
             return [
                Document(
                    page_content="[视频解析失败: 请确保服务器已安装 moviepy]",
                    metadata={"source": os.path.abspath(file_path), "error": True},
                )
            ]
        except Exception as e:
            logger.error(f"Video parsing error: {str(e)}")
            return [
                Document(
                    page_content="[视频解析失败: 请确保服务器已安装 FFmpeg]",
                    metadata={"source": os.path.abspath(file_path), "error": True},
                )
            ]

    def _ensure_dashscope_client(self) -> None:
        if self.dashscope_client is not None:
            return
        if not self.dashscope_api_key:
            return
        try:
            import dashscope
            from dashscope import MultiModalConversation

            dashscope.api_key = self.dashscope_api_key
            self.dashscope_client = MultiModalConversation
        except Exception:
            self.dashscope_client = None

    def _parse_image(self, file_path: str) -> List[Document]:
        try:
            self._ensure_dashscope_client()
            source = os.path.abspath(file_path)
            prompt = "请详细描述这张图片的内容，包括图表数据、文字信息和视觉细节。"
            if self.dashscope_client is None:
                message = "图片解析失败，请检查 API 配置"
                return [
                    Document(
                        page_content=message,
                        metadata={"source": source, "type": "image_error"},
                    )
                ]
            
            messages = [
                {
                    "role": "user",
                    "content": [
                        {"image": f"file://{source}"},
                        {"text": prompt},
                    ],
                }
            ]
            response = self.dashscope_client.call(
                model="qwen-vl-max",
                messages=messages,
            )
            output = getattr(response, "output", None) or {}
            choices = output.get("choices") or []
            if not choices:
                message = "图片解析失败，请检查 API 配置"
                return [
                    Document(
                        page_content=message,
                        metadata={"source": source, "type": "image_error"},
                    )
                ]
            message = choices[0].get("message") or {}
            contents = message.get("content") or []
            parts = []
            for item in contents:
                value = item.get("text")
                if value:
                    parts.append(value)
            text = "\n".join(parts).strip()
            if not text:
                message = "图片解析失败，请检查 API 配置"
                return [
                    Document(
                        page_content=message,
                        metadata={"source": source, "type": "image_error"},
                    )
                ]
            return [
                Document(
                    page_content=text,
                    metadata={"source": source, "type": "image_description"},
                )
            ]
        except Exception as e:
            logger.error(f"Image parsing error: {str(e)}")
            return [
                Document(
                    page_content=f"[图片解析失败: {str(e)}]",
                    metadata={"source": os.path.abspath(file_path), "type": "image_error"},
                )
            ]

    def _parse_txt(self, file_path: str) -> List[Document]:
        try:
            with open(file_path, "r", encoding="utf-8") as handle:
                text = handle.read()
            text = text.strip()
            if not text:
                return []
            source = os.path.abspath(file_path)
            return [
                Document(
                    page_content=text,
                    metadata={"source": source},
                )
            ]
        except Exception as e:
            logger.error(f"Txt parsing error: {str(e)}")
            return [
                Document(
                    page_content=f"[文本解析失败: {str(e)}]",
                    metadata={"source": os.path.abspath(file_path), "error": True},
                )
            ]
